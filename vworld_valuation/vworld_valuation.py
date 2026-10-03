"""브이월드 오픈API로 필지별 지목·면적·개별공시지가를 조회해 평가액 엑셀을 만든다.

사용법:
    pip install openpyxl requests
    python vworld_valuation.py --key 발급받은인증키 [--domain 등록한서비스URL]
    python vworld_valuation.py --offline   # API 호출 없이 빈 양식만 생성

PNU(19자리) = 법정동코드(10) + 산구분(1: 일반, 2: 산) + 본번(4) + 부번(4)
"""
import argparse
import sys
import time

import requests
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ADDRESS = "전북특별자치도 익산시 망성면 무형리"
# 법정동코드: 익산시 망성면 무형리 = 4514038026 (전라북도 시절).
# 2024-01-18 전북특별자치도 출범으로 시도코드가 45 -> 52 로 변경됨.
BJD_CODE = "5214038026"
BJD_CODE_OLD = "4514038026"
STDR_YEAR = "2026"
JIBUNS = [
    "458", "474", "474-3", "474-7", "475", "476", "477-12", "477-14",
    "481-16", "481-17", "490-14", "490-15", "490-16", "692-3", "694-1",
    "710-4", "710-6", "710-7", "745", "750", "757", "758", "759-1", "759-2",
    "760", "1816-3", "1820-1",
]

LAND_URL = "https://api.vworld.kr/ned/data/ladfrlList"              # 토지임야목록(지목·면적)
PRICE_URL = "https://api.vworld.kr/ned/data/getIndvdLandPriceAttr"  # 개별공시지가


def make_pnu(bjd, jibun, san=False):
    main, _, sub = jibun.partition("-")
    return f"{bjd}{'2' if san else '1'}{int(main):04d}{int(sub or 0):04d}"


def find_rows(obj):
    """응답 JSON 안에서 pnu 를 가진 dict 목록을 찾아낸다(응답 래핑 구조 차이 흡수)."""
    if isinstance(obj, list):
        if obj and all(isinstance(x, dict) for x in obj) and any("pnu" in x for x in obj):
            return obj
        for x in obj:
            r = find_rows(x)
            if r:
                return r
    elif isinstance(obj, dict):
        if "pnu" in obj:
            return [obj]
        for v in obj.values():
            r = find_rows(v)
            if r:
                return r
    return []


def call(url, params):
    for attempt in range(3):
        try:
            resp = requests.get(url, params=params, timeout=20)
            resp.raise_for_status()
            return resp.json()
        except Exception as e:  # noqa: BLE001
            if attempt == 2:
                print(f"  ! 호출 실패 {url} {params.get('pnu')}: {e}", file=sys.stderr)
                return {}
            time.sleep(1 + attempt)


def query(key, domain, pnu):
    base = {"key": key, "pnu": pnu, "format": "json", "numOfRows": "100", "pageNo": "1"}
    if domain:
        base["domain"] = domain
    land = find_rows(call(LAND_URL, base))
    prices = find_rows(call(PRICE_URL, {**base, "stdrYear": STDR_YEAR}))
    price = next((p for p in prices if str(p.get("stdrYear")) == STDR_YEAR), None)
    return (land[0] if land else None), price


def lookup(key, domain, jibun):
    """일반 필지 -> 산 필지 -> 구 법정동코드(45) 순으로 시도."""
    tried = []
    for bjd in (BJD_CODE, BJD_CODE_OLD):
        for san in (False, True):
            pnu = make_pnu(bjd, jibun, san)
            tried.append(pnu)
            land, price = query(key, domain, pnu)
            if land or price:
                note = []
                if san:
                    note.append("산 필지")
                if bjd != BJD_CODE:
                    note.append("구 법정동코드로 조회됨")
                if not price:
                    note.append(f"{STDR_YEAR}년 공시지가 없음")
                if not land:
                    note.append("토지대장 정보 없음")
                return pnu, land, price, ", ".join(note)
    return make_pnu(BJD_CODE, jibun), None, None, "조회 결과 없음(" + "/".join(tried) + ")"


def to_num(v):
    try:
        return float(str(v).replace(",", ""))
    except (TypeError, ValueError):
        return None


def build_excel(rows, path, offline):
    wb = Workbook()
    ws = wb.active
    ws.title = "평가액"
    font = "Arial"
    thin = Side(style="thin", color="999999")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill = PatternFill("solid", fgColor="D9E1F2")
    input_fill = PatternFill("solid", fgColor="FFFF00")
    tot_fill = PatternFill("solid", fgColor="F2F2F2")

    ws["A1"] = f"{ADDRESS} 27필지 평가액 ({STDR_YEAR}년 개별공시지가 기준)"
    ws["A1"].font = Font(name=font, bold=True, size=13)
    ws["A2"] = ("출처: 브이월드 오픈API 토지임야목록(ladfrlList)·개별공시지가(getIndvdLandPriceAttr). "
                "평가액 = 면적(㎡) × 개별공시지가(원/㎡).")
    if offline:
        ws["A2"].value += " ※ 노란 칸(지목·면적·공시지가)은 API 미조회 상태 — 값을 입력하면 평가액·합계가 자동 계산됨."
    ws["A2"].font = Font(name=font, size=9, color="595959")

    headers = ["No", "소재지", "지번", "PNU", "지목", "면적(㎡)", "개별공시지가(원/㎡)", "평가액(원)", "비고"]
    hr = 4
    for c, h in enumerate(headers, 1):
        cell = ws.cell(hr, c, h)
        cell.font = Font(name=font, bold=True)
        cell.fill = hdr_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    first = hr + 1
    for i, r in enumerate(rows):
        row = first + i
        vals = [i + 1, ADDRESS, r["jibun"], r["pnu"], r["jimok"], r["area"], r["price"],
                f"=IF(OR(F{row}=\"\",G{row}=\"\"),\"\",F{row}*G{row})", r["note"]]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row, c, v)
            cell.font = Font(name=font, color="0000FF" if c in (5, 6, 7) else "000000")
            cell.border = border
            if offline and c in (5, 6, 7):
                cell.fill = input_fill
        ws.cell(row, 4).number_format = "@"
        ws.cell(row, 6).number_format = "#,##0.00"
        ws.cell(row, 7).number_format = "#,##0"
        ws.cell(row, 8).number_format = "#,##0"
        for c in (1, 3, 4, 5):
            ws.cell(row, c).alignment = Alignment(horizontal="center")

    last = first + len(rows) - 1
    tot = last + 1
    ws.cell(tot, 1, "합계")
    ws.merge_cells(start_row=tot, start_column=1, end_row=tot, end_column=5)
    ws.cell(tot, 6, f"=SUM(F{first}:F{last})").number_format = "#,##0.00"
    ws.cell(tot, 8, f"=SUM(H{first}:H{last})").number_format = "#,##0"
    for c in range(1, 10):
        cell = ws.cell(tot, c)
        cell.font = Font(name=font, bold=True)
        cell.fill = tot_fill
        cell.border = border
    ws.cell(tot, 1).alignment = Alignment(horizontal="center")

    ws.cell(first, 4).comment = Comment(
        f"PNU = 법정동코드 {BJD_CODE}(익산시 망성면 무형리, 전북특별자치도 52) + 산구분 1(일반) "
        "+ 본번 4자리 + 부번 4자리", "script")

    for col, w in zip("ABCDEFGHI", [5, 34, 9, 22, 8, 13, 19, 18, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(first, 1)
    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--key", help="브이월드 인증키")
    ap.add_argument("--domain", default="", help="인증키 발급 시 등록한 서비스 URL(필요한 경우)")
    ap.add_argument("--offline", action="store_true", help="API 호출 없이 빈 양식 생성")
    ap.add_argument("-o", "--output", default="무형리_27필지_평가액.xlsx")
    a = ap.parse_args()
    if not a.offline and not a.key:
        ap.error("--key 가 필요합니다 (또는 --offline)")

    rows = []
    for j in JIBUNS:
        if a.offline:
            rows.append(dict(jibun=j, pnu=make_pnu(BJD_CODE, j), jimok=None, area=None, price=None, note=""))
            continue
        pnu, land, price, note = lookup(a.key, a.domain, j)
        rows.append(dict(
            jibun=j, pnu=pnu,
            jimok=(land or {}).get("lndcgrCodeNm"),
            area=to_num((land or {}).get("lndpclAr")),
            price=to_num((price or {}).get("pblntfPclnd")),
            note=note,
        ))
        r = rows[-1]
        print(f"{j:>8} {pnu} {r['jimok']} {r['area']} {r['price']} {note}")
    build_excel(rows, a.output, a.offline)
    print(f"저장: {a.output}")


if __name__ == "__main__":
    main()
